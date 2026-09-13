import os
import sys
import uuid
import io
import time
import asyncio
import hashlib
import csv
import docx
import boto3
import psycopg2
import json 
from contextlib import asynccontextmanager, contextmanager
from psycopg2 import pool
from psycopg2.extras import execute_values 
from fastapi import FastAPI, UploadFile, File, HTTPException, BackgroundTasks, Depends, Request, status
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
import jwt
from jwt import PyJWKClient
from pypdf import PdfReader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from fastembed import TextEmbedding
from pydantic import BaseModel
from typing import Optional
from groq import AsyncGroq
from loguru import logger
from slowapi import Limiter, _rate_limit_exceeded_handler
from slowapi.util import get_remote_address
from slowapi.errors import RateLimitExceeded

load_dotenv()

# --- Structured Logging Setup ---
logger.remove()
LOG_JSON = os.getenv("LOG_JSON", "false").lower() == "true"
if LOG_JSON:
    logger.add(sys.stdout, serialize=True)
else:
    logger.add(
        sys.stdout,
        format="<green>{time:YYYY-MM-DD HH:mm:ss.SSS}</green> | <level>{level: <8}</level> | <cyan>{name}</cyan>:<cyan>{function}</cyan>:<cyan>{line}</cyan> - <level>{message}</level>",
        level=os.getenv("LOG_LEVEL", "INFO")
    )

# --- Database Connection Pooling & Initialization ---
db_pool: Optional[pool.ThreadedConnectionPool] = None

def get_db_pool() -> pool.ThreadedConnectionPool:
    global db_pool
    if db_pool is None:
        db_url = os.getenv("DATABASE_URL")
        if not db_url:
            raise RuntimeError("DATABASE_URL environment variable is not set")
        db_pool = pool.ThreadedConnectionPool(
            minconn=1,
            maxconn=15,
            dsn=db_url
        )
    return db_pool

@contextmanager
def get_db_connection():
    """Context manager yielding a pooled database connection and returning it on exit."""
    p = get_db_pool()
    conn = p.getconn()
    try:
        yield conn
    finally:
        p.putconn(conn)

def init_db():
    try:
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                cur.execute("""
                    ALTER TABLE documents 
                    ADD COLUMN IF NOT EXISTS status VARCHAR(50) DEFAULT 'ready',
                    ADD COLUMN IF NOT EXISTS error_message TEXT,
                    ADD COLUMN IF NOT EXISTS file_hash VARCHAR(64);
                    CREATE INDEX IF NOT EXISTS idx_documents_user_hash ON documents(user_id, file_hash);
                    
                    ALTER TABLE document_chunks
                    ADD COLUMN IF NOT EXISTS page_number INTEGER;
                """)
                conn.commit()
        logger.info("Database schema migration verified.")
    except Exception as e:
        logger.warning(f"Schema migration check failed/skipped: {e}")

@asynccontextmanager
async def lifespan(app: FastAPI):
    # Startup: Verify database schema
    init_db()
    yield
    # Shutdown: Close database pool
    global db_pool
    if db_pool:
        db_pool.closeall()
        logger.info("Database connection pool closed.")

# --- Rate Limiter & FastAPI App Setup ---
limiter = Limiter(key_func=get_remote_address, default_limits=["120/minute"])
app = FastAPI(title="Pagewise API", lifespan=lifespan)
app.state.limiter = limiter
app.add_exception_handler(RateLimitExceeded, _rate_limit_exceeded_handler)

class SearchQuery(BaseModel):
    query: str
    document_id: Optional[str] = None 
    top_k: int = 3 
    user_id: Optional[str] = None 

# Setup Clerk JWKS JWT verification
security = HTTPBearer()
CLERK_JWKS_URL = os.getenv("CLERK_JWKS_URL")
jwks_client = PyJWKClient(CLERK_JWKS_URL) if CLERK_JWKS_URL else None

async def get_current_user(credentials: HTTPAuthorizationCredentials = Depends(security)) -> str:
    token = credentials.credentials
    if not token:
        logger.warning("Authentication failed: Missing bearer token")
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Missing bearer authentication token",
            headers={"WWW-Authenticate": "Bearer"},
        )
    if not jwks_client:
        logger.error("Authentication failed: Clerk JWKS URL not configured")
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail="Clerk JWKS URL is not configured on the server",
        )
    try:
        signing_key = jwks_client.get_signing_key_from_jwt(token)
        payload = jwt.decode(
            token,
            signing_key.key,
            algorithms=["RS256"],
            options={"verify_exp": True, "verify_aud": False},
        )
        user_id = payload.get("sub")
        if not user_id:
            logger.warning("Authentication failed: Token missing sub claim")
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Token does not contain a valid user identifier (sub)",
                headers={"WWW-Authenticate": "Bearer"},
            )
        return user_id
    except jwt.PyJWTError as e:
        logger.warning(f"Authentication token validation failed: {e}")
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail=f"Invalid or expired authentication token: {str(e)}",
            headers={"WWW-Authenticate": "Bearer"},
        )

raw_origins = os.getenv("ALLOWED_ORIGINS", "")
origins = [origin.strip() for origin in raw_origins.split(",") if origin.strip()]
if not origins:
    origins = [
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def health_check():
    return {"status": "ok", "message": "Pagewise API is running"}

logger.info("Loading embedding model (all-MiniLM-L6-v2 via FastEmbed)...")
model = TextEmbedding(model_name="sentence-transformers/all-MiniLM-L6-v2")
logger.info("Embedding model loaded successfully!")

logger.info("Initializing Groq Client...")
groq_client = AsyncGroq(api_key=os.getenv("GROQ_API_KEY"), timeout=15.0)
logger.info("Groq Client ready!")

s3 = boto3.client(
    "s3",
    endpoint_url=os.getenv("SUPABASE_ENDPOINT_URL"),
    aws_access_key_id=os.getenv("SUPABASE_ACCESS_KEY_ID"),
    aws_secret_access_key=os.getenv("SUPABASE_SECRET_ACCESS_KEY"),
    region_name="us-east-1", 
)
BUCKET_NAME = os.getenv("SUPABASE_BUCKET_NAME")

# --- External API Resilience Helpers ---
GROQ_MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")
FALLBACK_GROQ_MODELS = [GROQ_MODEL, "llama-3.3-70b-versatile"]

async def create_groq_stream_with_retry(messages: list, max_retries: int = 2):
    """Initializes Groq streaming completion with exponential backoff and automatic model fallback."""
    last_err = None
    # Deduplicate while preserving order
    candidate_models = list(dict.fromkeys(FALLBACK_GROQ_MODELS))
    
    for model_name in candidate_models:
        for attempt in range(max_retries + 1):
            try:
                stream = await groq_client.chat.completions.create(
                    messages=messages,
                    model=model_name,
                    temperature=0.2,
                    stream=True
                )
                return stream
            except Exception as e:
                last_err = e
                err_str = str(e).lower()
                # If model is deprecated or not found for this key, try next candidate model immediately
                if "not found" in err_str or "does not exist" in err_str or "model_not_found" in err_str:
                    logger.warning(f"Groq model '{model_name}' not available for account, attempting fallback model...")
                    break
                
                logger.warning(f"Groq stream init failed on '{model_name}' (attempt {attempt + 1}/{max_retries + 1}): {e}")
                if attempt < max_retries:
                    await asyncio.sleep(0.5 * (2 ** attempt))

    if last_err:
        raise last_err

def extract_pdf(file_bytes: bytes) -> list[tuple[str, Optional[int]]]:
    """Extracts text page-by-page from PDF files using pypdf."""
    reader = PdfReader(io.BytesIO(file_bytes))
    results = []
    for page_idx, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        page_text = page_text.replace('\x00', '').strip()
        if page_text:
            results.append((page_text, page_idx))
    return results

def extract_docx(file_bytes: bytes) -> list[tuple[str, Optional[int]]]:
    """Extracts text from DOCX files including paragraphs and tables."""
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    tables_text = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                tables_text.append(" | ".join(row_cells))
                
    full_text = "\n\n".join(paragraphs_text + tables_text).strip()
    return [(full_text, None)] if full_text else []

def extract_csv(file_bytes: bytes) -> list[tuple[str, Optional[int]]]:
    """Extracts tabular data from CSV files and formats rows into readable key-value text."""
    text_content = file_bytes.decode("utf-8", errors="replace")
    csv_file = io.StringIO(text_content)
    reader = csv.reader(csv_file)
    
    rows = [row for row in reader if any(field.strip() for field in row)]
    if not rows:
        return []
        
    header = rows[0]
    has_header = len(rows) > 1 and any(h.strip() for h in header)
    
    formatted_lines = []
    if has_header:
        headers = [h.strip() or f"col_{i+1}" for i, h in enumerate(header)]
        formatted_lines.append("Columns: " + ", ".join(headers))
        for row_idx, row in enumerate(rows[1:], start=1):
            row_data = []
            for col_idx, cell in enumerate(row):
                header_name = headers[col_idx] if col_idx < len(headers) else f"col_{col_idx+1}"
                row_data.append(f"{header_name}: {cell.strip()}")
            formatted_lines.append(f"Row {row_idx}: " + ", ".join(row_data))
    else:
        for row_idx, row in enumerate(rows, start=1):
            formatted_lines.append(f"Row {row_idx}: " + ", ".join(cell.strip() for cell in row))
            
    full_text = "\n".join(formatted_lines).strip()
    return [(full_text, None)] if full_text else []

def extract_text(file_bytes: bytes, file_ext: str) -> list[tuple[str, Optional[int]]]:
    """Pluggable dispatcher routing file bytes to format-specific text extractors."""
    ext = file_ext.lower().lstrip(".")
    if ext == "pdf":
        return extract_pdf(file_bytes)
    elif ext == "docx":
        return extract_docx(file_bytes)
    elif ext == "csv":
        return extract_csv(file_bytes)
    else:
        raise ValueError(f"Unsupported file format: .{ext}")

def process_document_background(file_bytes: bytes, doc_id: str, file_ext: str):
    try:
        start_time = time.perf_counter()
        logger.info(f"[{doc_id}] Starting background AI processing for format '{file_ext}'...")
        
        extracted_sections = extract_text(file_bytes, file_ext)
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
        
        all_chunks = []
        for section_text, page_idx in extracted_sections:
            if not section_text.strip():
                continue
            section_chunks = splitter.split_text(section_text)
            for chunk in section_chunks:
                all_chunks.append((chunk, page_idx))
        
        if not all_chunks:
            logger.warning(f"[{doc_id}] No extractable text found in .{file_ext} file.")
            with get_db_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE documents SET status = %s, error_message = %s WHERE id = %s",
                        ("empty", f"No extractable text found in .{file_ext} file", doc_id)
                    )
                    conn.commit()
            return

        logger.info(f"[{doc_id}] Encoding {len(all_chunks)} chunks...")
        chunk_texts = [c[0] for c in all_chunks]
        embeddings = [e.tolist() for e in model.embed(chunk_texts)]

        records = [
            (str(uuid.uuid4()), doc_id, i, chunk_text, embedding, page_num)
            for i, ((chunk_text, page_num), embedding) in enumerate(zip(all_chunks, embeddings))
        ]

        with get_db_connection() as conn:
            with conn.cursor() as cur:
                execute_values(
                    cur,
                    "INSERT INTO document_chunks (id, document_id, chunk_index, chunk_text, embedding, page_number) VALUES %s",
                    records
                )
                
                cur.execute(
                    "UPDATE documents SET status = %s, error_message = NULL WHERE id = %s",
                    ("ready", doc_id)
                )
                conn.commit()

        total_elapsed = (time.perf_counter() - start_time) * 1000
        logger.info(f"[{doc_id}] Ingestion complete: {len(all_chunks)} chunks saved in {total_elapsed:.1f}ms")

    except Exception as e:
        logger.error(f"[{doc_id}] Background processing failed: {e}")
        try:
            with get_db_connection() as conn:
                with conn.cursor() as cur:
                    cur.execute(
                        "UPDATE documents SET status = %s, error_message = %s WHERE id = %s",
                        ("failed", str(e), doc_id)
                    )
                    conn.commit()
        except Exception as db_err:
            logger.error(f"Failed to record failure status for {doc_id}: {db_err}")

@app.post("/upload")
@limiter.limit("5/minute")
async def upload_document(
    request: Request,
    background_tasks: BackgroundTasks, 
    file: UploadFile = File(...), 
    user_id: str = Depends(get_current_user)
):
    upload_start = time.perf_counter()
    logger.info(f"Upload request received: filename='{file.filename}', user_id='{user_id}'")
    
    filename = file.filename or ""
    file_ext = os.path.splitext(filename)[1].lower().lstrip(".")
    if file_ext not in ["pdf", "docx", "csv"]:
        raise HTTPException(status_code=400, detail="Only PDF, DOCX, and CSV files are allowed")

    doc_id = str(uuid.uuid4())
    s3_key = f"{user_id}/{doc_id}_{filename}"

    try:
        file_bytes = await file.read()
        file_hash = hashlib.sha256(file_bytes).hexdigest()

        # Deduplication check: return existing document if identical hash exists for this user
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id, filename, status FROM documents WHERE user_id = %s AND file_hash = %s",
                    (user_id, file_hash)
                )
                existing = cur.fetchone()
                if existing:
                    logger.info(f"Document deduplicated: hash={file_hash[:8]} matched existing doc_id={existing[0]}")
                    return {
                        "message": f"Document '{existing[1]}' has already been uploaded.",
                        "document_id": existing[0],
                        "status": existing[2] or "ready",
                        "deduplicated": True
                    }

        content_type_map = {
            "pdf": "application/pdf",
            "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "csv": "text/csv",
        }
        content_type = content_type_map.get(file_ext, "application/octet-stream")

        s3.upload_fileobj(
            io.BytesIO(file_bytes), 
            BUCKET_NAME, 
            s3_key,
            ExtraArgs={"ContentType": content_type}
        )
        logger.info(f"Uploaded file to S3: bucket='{BUCKET_NAME}', key='{s3_key}'")

        with get_db_connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO documents (id, user_id, filename, s3_key, status, file_hash) VALUES (%s, %s, %s, %s, %s, %s)",
                    (doc_id, user_id, filename, s3_key, "processing", file_hash)
                )
                conn.commit()

        background_tasks.add_task(process_document_background, file_bytes, doc_id, file_ext)

        elapsed_ms = (time.perf_counter() - upload_start) * 1000
        logger.info(f"Upload completed in {elapsed_ms:.1f}ms: doc_id='{doc_id}' queued for background ingestion")
        return {
            "message": "File uploaded! AI is processing your document in the background.",
            "document_id": doc_id,
            "status": "processing",
            "deduplicated": False
        }

    except Exception as e:
        logger.error(f"Upload error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/documents")
@limiter.limit("60/minute")
async def list_documents(request: Request, user_id: str = Depends(get_current_user)):
    try:
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT id, filename, s3_key, COALESCE(status, 'ready') as status, error_message FROM documents WHERE user_id = %s ORDER BY filename ASC",
                    (user_id,)
                )
                rows = cur.fetchall()
        
        documents = [
            {
                "id": row[0],
                "filename": row[1],
                "s3_key": row[2],
                "status": row[3],
                "error_message": row[4]
            }
            for row in rows
        ]
        return documents
    except Exception as e:
        logger.error(f"List documents error for user '{user_id}': {e}")
        raise HTTPException(status_code=500, detail=str(e))

@app.delete("/documents/{document_id}")
async def delete_document(document_id: str, user_id: str = Depends(get_current_user)):
    try:
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT s3_key FROM documents WHERE id = %s AND user_id = %s", (document_id, user_id))
                row = cur.fetchone()
                if not row:
                    raise HTTPException(status_code=404, detail="Document not found or unauthorized")
                
                s3_key = row[0]
                
                try:
                    s3.delete_object(Bucket=BUCKET_NAME, Key=s3_key)
                    logger.info(f"Deleted S3 object: {s3_key}")
                except Exception as s3_err:
                    logger.warning(f"Failed to delete S3 object {s3_key}: {s3_err}")
                
                cur.execute("DELETE FROM chat_messages WHERE document_id = %s AND user_id = %s", (document_id, user_id))
                cur.execute("DELETE FROM document_chunks WHERE document_id = %s", (document_id,))
                cur.execute("DELETE FROM documents WHERE id = %s AND user_id = %s", (document_id, user_id))
                conn.commit()
        
        logger.info(f"Document {document_id} and associated chunks/messages successfully deleted for user {user_id}")
        return {"message": "Document successfully deleted"}
    except HTTPException as he:
        raise he
    except Exception as e:
        logger.error(f"Delete document error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/chat/history")
async def get_chat_history(document_id: Optional[str] = None, user_id: str = Depends(get_current_user)):
    try:
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                if document_id:
                    cur.execute(
                        "SELECT role, content FROM chat_messages WHERE user_id = %s AND document_id = %s ORDER BY created_at ASC",
                        (user_id, document_id)
                    )
                else:
                    cur.execute(
                        "SELECT role, content FROM chat_messages WHERE user_id = %s AND document_id IS NULL ORDER BY created_at ASC",
                        (user_id,)
                    )
                rows = cur.fetchall()
        
        history = [{"role": row[0], "content": row[1]} for row in rows]
        return history
    except Exception as e:
        logger.error(f"Get chat history error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/chat")
@limiter.limit("20/minute")
async def chat_with_document(
    request: Request,
    search: SearchQuery, 
    user_id: str = Depends(get_current_user)
):
    try:
        pipeline_start = time.perf_counter()
        logger.info(f"Chat request started: query='{search.query}', doc_id='{search.document_id}', user_id='{user_id}'")

        emb_start = time.perf_counter()
        query_embeddings = await asyncio.to_thread(lambda: list(model.embed([search.query])))
        emb_elapsed_ms = (time.perf_counter() - emb_start) * 1000
        query_vector = query_embeddings[0].tolist()
        query_vector_str = "[" + ",".join(map(str, query_vector)) + "]"

        db_start = time.perf_counter()
        with get_db_connection() as conn:
            with conn.cursor() as cur:
                if search.document_id:
                    cur.execute("""
                        WITH vector_search AS (
                            SELECT dc.chunk_text, dc.page_number, d.filename,
                                   row_number() over (ORDER BY dc.embedding <=> %s::vector) as rank
                            FROM document_chunks dc
                            JOIN documents d ON d.id = dc.document_id
                            WHERE dc.document_id = %s
                            LIMIT 20
                        ),
                        keyword_search AS (
                            SELECT dc.chunk_text, dc.page_number, d.filename,
                                   row_number() over (ORDER BY ts_rank(dc.fts, websearch_to_tsquery('english', %s)) DESC) as rank
                            FROM document_chunks dc
                            JOIN documents d ON d.id = dc.document_id
                            WHERE dc.document_id = %s AND dc.fts @@ websearch_to_tsquery('english', %s)
                            LIMIT 20
                        )
                        SELECT COALESCE(v.chunk_text, k.chunk_text) AS chunk_text,
                               COALESCE(v.page_number, k.page_number) AS page_number,
                               COALESCE(v.filename, k.filename) AS filename,
                               COALESCE(1.0 / (v.rank + 60), 0.0) + COALESCE(1.0 / (k.rank + 60), 0.0) AS rrf_score
                        FROM vector_search v
                        FULL OUTER JOIN keyword_search k ON v.chunk_text = k.chunk_text
                        ORDER BY rrf_score DESC
                        LIMIT %s;
                    """, (query_vector_str, search.document_id, search.query, search.document_id, search.query, search.top_k))
                else:
                    cur.execute("""
                        WITH vector_search AS (
                            SELECT dc.chunk_text, dc.page_number, d.filename,
                                   row_number() over (ORDER BY dc.embedding <=> %s::vector) as rank
                            FROM document_chunks dc
                            JOIN documents d ON d.id = dc.document_id
                            WHERE d.user_id = %s
                            LIMIT 20
                        ),
                        keyword_search AS (
                            SELECT dc.chunk_text, dc.page_number, d.filename,
                                   row_number() over (ORDER BY ts_rank(dc.fts, websearch_to_tsquery('english', %s)) DESC) as rank
                            FROM document_chunks dc
                            JOIN documents d ON d.id = dc.document_id
                            WHERE d.user_id = %s AND dc.fts @@ websearch_to_tsquery('english', %s)
                            LIMIT 20
                        )
                        SELECT COALESCE(v.chunk_text, k.chunk_text) AS chunk_text,
                               COALESCE(v.page_number, k.page_number) AS page_number,
                               COALESCE(v.filename, k.filename) AS filename,
                               COALESCE(1.0 / (v.rank + 60), 0.0) + COALESCE(1.0 / (k.rank + 60), 0.0) AS rrf_score
                        FROM vector_search v
                        FULL OUTER JOIN keyword_search k ON v.chunk_text = k.chunk_text
                        ORDER BY rrf_score DESC
                        LIMIT %s;
                    """, (query_vector_str, user_id, search.query, user_id, search.query, search.top_k))

                final_chunks = cur.fetchall()

        db_elapsed_ms = (time.perf_counter() - db_start) * 1000

        # Build context for LLM prompt directly from RRF top-k chunks
        retrieved_context = "\n\n---\n\n".join([row[0] for row in final_chunks if row[0]])

        # Extract structured document citations for UI grounding
        doc_sources = []
        seen_doc_citations = set()
        for row in final_chunks:
            chunk_text, page_num, filename = row[0], row[1], row[2]
            if filename and chunk_text:
                citation_key = (filename, page_num)
                if citation_key not in seen_doc_citations:
                    seen_doc_citations.add(citation_key)
                    doc_sources.append({
                        "filename": filename,
                        "page_number": page_num,
                        "snippet": chunk_text[:160].strip() + ("..." if len(chunk_text) > 160 else "")
                    })

        logger.info(
            f"RAG telemetry | emb: {emb_elapsed_ms:.1f}ms | db: {db_elapsed_ms:.1f}ms "
            f"| chunks: {len(final_chunks)} | doc_sources: {len(doc_sources)}"
        )

        final_context = f"Local Document Context:\n{retrieved_context}"

        prompt = f"""
        You are a highly intelligent academic assistant. Answer the user's question directly and naturally.

        CRITICAL INSTRUCTIONS:
        - DO NOT start your response with phrases like "Based on the provided context". Just answer the question immediately.
        - DO NOT mention the term "Local Document Context" in your response.
        - DO NOT use any citation markers, brackets, or numbers (like [1], [2]) in your text response. Integrate the information seamlessly.
        - If the answer cannot be found, simply state: "I cannot answer this based on the available information."

        Context:
        {final_context}

        User Question:
        {search.query}
        """
        async def stream_groq_response():
            full_ai_response = ""
            
            with get_db_connection() as conn:
                with conn.cursor() as cur:
                    if search.document_id:
                        cur.execute("""
                            SELECT role, content FROM chat_messages 
                            WHERE document_id = %s AND user_id = %s
                            ORDER BY created_at DESC 
                            LIMIT 6
                        """, (search.document_id, user_id))
                    else:
                        cur.execute("""
                            SELECT role, content FROM chat_messages 
                            WHERE document_id IS NULL AND user_id = %s
                            ORDER BY created_at DESC 
                            LIMIT 6
                        """, (user_id,))
                    history_rows = cur.fetchall()

            history_rows.reverse()

            groq_messages = [
                {"role": "system", "content": "You are a helpful academic assistant."}
            ]
            
            for row in history_rows:
                groq_messages.append({"role": row[0], "content": row[1]})
                
            groq_messages.append({"role": "user", "content": prompt})

            stream = await create_groq_stream_with_retry(groq_messages, max_retries=2)
            
            async for chunk in stream:
                token = chunk.choices[0].delta.content
                if token:
                    full_ai_response += token
                    yield token

            # Package document citations
            sources_payload = {
                "documents": doc_sources
            }
            if doc_sources:
                yield f"\n\n<<<SOURCES>>>{json.dumps(sources_payload)}"
                
            try:
                with get_db_connection() as conn:
                    with conn.cursor() as cur:
                        user_msg_id = str(uuid.uuid4())
                        ai_msg_id = str(uuid.uuid4())
                        
                        cur.execute(
                            "INSERT INTO chat_messages (id, document_id, role, content, user_id) VALUES (%s, %s, %s, %s, %s)",
                            (user_msg_id, search.document_id, "user", search.query, user_id)
                        )
                        
                        cur.execute(
                            "INSERT INTO chat_messages (id, document_id, role, content, user_id) VALUES (%s, %s, %s, %s, %s)",
                            (ai_msg_id, search.document_id, "assistant", full_ai_response, user_id)
                        )
                        conn.commit()
                total_pipeline_ms = (time.perf_counter() - pipeline_start) * 1000
                logger.info(f"Chat stream completed & persisted in {total_pipeline_ms:.1f}ms")
                
            except Exception as db_error:
                logger.error(f"Database error saving history: {str(db_error)}")

        return StreamingResponse(stream_groq_response(), media_type="text/plain")

    except Exception as e:
        logger.error(f"Generation error: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)