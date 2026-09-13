import io
import hashlib
import docx
import pytest
from unittest.mock import patch, MagicMock
from main import (
    extract_pdf,
    extract_docx,
    extract_csv,
    extract_text,
    process_document_background,
)

def test_upload_unsupported_format_rejected(auth_client):
    response = auth_client.post(
        "/upload",
        files={"file": ("test.txt", b"Hello world", "text/plain")}
    )
    assert response.status_code == 400
    assert "Only PDF, DOCX, and CSV files are allowed" in response.json()["detail"]

def test_upload_image_rejected(auth_client):
    response = auth_client.post(
        "/upload",
        files={"file": ("photo.jpg", b"fake_image_bytes", "image/jpeg")}
    )
    assert response.status_code == 400
    assert "Only PDF, DOCX, and CSV files are allowed" in response.json()["detail"]

def test_upload_deduplication(auth_client):
    fake_pdf_content = b"%PDF-1.4 Fake PDF Content for deduplication test"
    file_hash = hashlib.sha256(fake_pdf_content).hexdigest()

    mock_db = MagicMock()
    mock_cursor = MagicMock()
    # Simulate existing document found
    mock_cursor.fetchone.return_value = ("existing-doc-id-123", "test.pdf", "ready")
    mock_db.cursor.return_value.__enter__.return_value = mock_cursor

    with patch("main.get_db_connection") as mock_conn:
        mock_conn.return_value.__enter__.return_value = mock_db
        response = auth_client.post(
            "/upload",
            files={"file": ("test.pdf", fake_pdf_content, "application/pdf")}
        )
        assert response.status_code == 200
        data = response.json()
        assert data["deduplicated"] is True
        assert data["document_id"] == "existing-doc-id-123"

def test_extract_pdf_normal():
    mock_pdf_reader = MagicMock()
    page1 = MagicMock()
    page1.extract_text.return_value = "Page 1 content about neural networks."
    page2 = MagicMock()
    page2.extract_text.return_value = "Page 2 content about transformers."
    mock_pdf_reader.pages = [page1, page2]

    with patch("main.PdfReader", return_value=mock_pdf_reader):
        sections = extract_pdf(b"fake_pdf_bytes")
        assert len(sections) == 2
        assert sections[0] == ("Page 1 content about neural networks.", 1)
        assert sections[1] == ("Page 2 content about transformers.", 2)

def test_extract_pdf_empty():
    mock_pdf_reader = MagicMock()
    page1 = MagicMock()
    page1.extract_text.return_value = ""
    mock_pdf_reader.pages = [page1]

    with patch("main.PdfReader", return_value=mock_pdf_reader):
        sections = extract_pdf(b"empty_pdf_bytes")
        assert sections == []

def test_extract_docx_normal():
    doc = docx.Document()
    doc.add_paragraph("Introduction to RAG pipelines.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Col A"
    table.rows[0].cells[1].text = "Col B"
    
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    sections = extract_docx(docx_bytes)
    assert len(sections) == 1
    assert "Introduction to RAG pipelines." in sections[0][0]
    assert "Col A | Col B" in sections[0][0]
    assert sections[0][1] is None  # Page number is None for DOCX

def test_extract_docx_empty():
    doc = docx.Document()
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    sections = extract_docx(docx_bytes)
    assert sections == []

def test_extract_csv_normal():
    csv_bytes = b"Name,Role,Department\nAlice,Engineer,AI\nBob,Researcher,ML"
    sections = extract_csv(csv_bytes)
    assert len(sections) == 1
    text, page_num = sections[0]
    assert page_num is None
    assert "Columns: Name, Role, Department" in text
    assert "Row 1: Name: Alice, Role: Engineer, Department: AI" in text
    assert "Row 2: Name: Bob, Role: Researcher, Department: ML" in text

def test_extract_csv_empty():
    csv_bytes = b"   \n,,\n  "
    sections = extract_csv(csv_bytes)
    assert sections == []

def test_extract_text_dispatcher():
    # PDF dispatch
    mock_pdf_reader = MagicMock()
    page1 = MagicMock()
    page1.extract_text.return_value = "PDF text"
    mock_pdf_reader.pages = [page1]
    with patch("main.PdfReader", return_value=mock_pdf_reader):
        res_pdf = extract_text(b"pdf_bytes", ".pdf")
        assert res_pdf == [("PDF text", 1)]

    # CSV dispatch
    res_csv = extract_text(b"col1,col2\nval1,val2", "csv")
    assert len(res_csv) == 1
    assert res_csv[0][1] is None

    # Unsupported format
    with pytest.raises(ValueError) as exc:
        extract_text(b"raw", "exe")
    assert "Unsupported file format" in str(exc.value)

def test_process_document_background_empty():
    mock_db = MagicMock()
    mock_cursor = MagicMock()
    mock_db.cursor.return_value.__enter__.return_value = mock_cursor

    mock_pdf_reader = MagicMock()
    mock_page = MagicMock()
    mock_page.extract_text.return_value = ""
    mock_pdf_reader.pages = [mock_page]

    with patch("main.PdfReader", return_value=mock_pdf_reader), \
         patch("main.get_db_connection") as mock_conn:
        mock_conn.return_value.__enter__.return_value = mock_db

        process_document_background(b"empty_pdf_bytes", "empty_doc_id", "pdf")

        # Verify DB update status set to empty
        mock_cursor.execute.assert_called_once()
        args = mock_cursor.execute.call_args[0]
        assert "empty" in args[1]
